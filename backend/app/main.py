import io
import os
from functools import lru_cache

from fastapi import FastAPI, File, Form, HTTPException, UploadFile, Depends
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from pydantic import BaseModel

import pdfplumber
from docx import Document
import spacy
from sentence_transformers import SentenceTransformer, util
from dotenv import load_dotenv
from thefuzz import fuzz

from .scoring import analyze_with_rules
from .roadmap import generate_roadmap
from .database import Base, engine, get_db
from .models import Analysis, User
from .auth import hash_password, verify_password

load_dotenv(override=True)


# ---------- Request models ----------
class SignupRequest(BaseModel):
    username: str
    email: str
    password: str


class LoginRequest(BaseModel):
    username: str
    password: str


app = FastAPI(title="AI Skill Gap Analyzer API", version="0.9.2")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# Create tables on startup (no-op if they already exist)
Base.metadata.create_all(bind=engine)


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


# ---------- Models (loaded once at startup) ----------

try:
    nlp = spacy.load("en_core_web_sm")
except OSError:
    nlp = spacy.blank("en")


LOCAL_MODEL_DIR = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "models", "all-MiniLM-L6-v2")
)

_embedder = None


def get_embedder():
    global _embedder
    if _embedder is None:
        if os.path.isdir(LOCAL_MODEL_DIR):
            print(f"[embedder] loading from local dir: {LOCAL_MODEL_DIR}")
            _embedder = SentenceTransformer(LOCAL_MODEL_DIR)
        else:
            print("[embedder] local dir not found, falling back to hub name")
            _embedder = SentenceTransformer("all-MiniLM-L6-v2")
    return _embedder


@lru_cache(maxsize=2048)
def embed_skill(skill: str):
    model = get_embedder()
    return model.encode(skill, normalize_embeddings=True, convert_to_tensor=True)


# ---------- Skill dictionary ----------
SKILL_LIST = [
    "Python", "Java", "JavaScript", "TypeScript", "C", "C++", "C#", "Go", "Rust",
    "Ruby", "PHP", "Swift", "Kotlin", "Scala", "R", "MATLAB", "Perl", "Dart",
    "Bash", "Shell", "SQL", "NoSQL", "HTML", "CSS", "Sass", "SCSS",

    "React", "React Native", "Angular", "Vue", "Vue.js", "Svelte", "Next.js",
    "Nuxt.js", "Redux", "MobX", "Tailwind CSS", "Bootstrap", "Material UI",
    "jQuery", "Webpack", "Vite", "Babel", "ESLint",

    "Node.js", "Express", "NestJS", "Django", "Flask", "FastAPI", "Spring",
    "Spring Boot", "Laravel", "Ruby on Rails", "ASP.NET", ".NET", "GraphQL",
    "REST", "REST API", "gRPC", "WebSockets", "Microservices",

    "PostgreSQL", "MySQL", "MongoDB", "SQLite", "Redis", "Cassandra",
    "DynamoDB", "Firebase", "Oracle", "SQL Server", "Elasticsearch", "Neo4j",
    "Snowflake", "BigQuery", "Redshift",

    "AWS", "Azure", "GCP", "Google Cloud", "Docker", "Kubernetes", "Terraform",
    "Ansible", "Jenkins", "GitHub Actions", "GitLab CI", "CircleCI", "CI/CD",
    "Linux", "Nginx", "Apache", "Heroku", "Vercel", "Netlify", "Cloudflare",
    "Serverless", "Lambda", "EC2", "S3",

    "Machine Learning", "Deep Learning", "Data Science", "Data Analysis",
    "Data Engineering", "NLP", "Natural Language Processing",
    "Computer Vision", "Reinforcement Learning", "TensorFlow", "PyTorch",
    "Keras", "Scikit-learn", "XGBoost", "LightGBM", "Pandas", "NumPy",
    "Matplotlib", "Seaborn", "Plotly", "Jupyter", "OpenCV", "Hugging Face",
    "Transformers", "LLM", "Generative AI", "Prompt Engineering",
    "Statistics", "Probability", "A/B Testing", "Feature Engineering",

    "Hadoop", "Spark", "PySpark", "Kafka", "Airflow", "Databricks", "ETL",

    "Android", "iOS", "Flutter", "Xamarin",

    "Git", "GitHub", "GitLab", "Bitbucket", "Jira", "Confluence", "Slack",
    "VS Code", "IntelliJ", "Postman", "Swagger", "Figma", "Notion",
    "Agile", "Scrum", "Kanban", "TDD", "Unit Testing", "Pytest", "Jest",
    "Selenium", "Cypress", "Playwright", "OOP", "Design Patterns",
    "Data Structures", "Algorithms", "System Design", "Operating Systems",
    "Computer Networks", "DBMS",

    "Cybersecurity", "Penetration Testing", "OWASP", "OAuth", "JWT",

    "Communication", "Teamwork", "Leadership", "Problem Solving",
    "Project Management", "Time Management",
]


SKILL_ALIASES = {
    "nodejs": "Node.js",
    "node": "Node.js",
    "node js": "Node.js",
    "reactjs": "React",
    "react.js": "React",
    "vuejs": "Vue",
    "vue js": "Vue",
    "nextjs": "Next.js",
    "next js": "Next.js",
    "expressjs": "Express",
    "express.js": "Express",
    "postgres": "PostgreSQL",
    "mongo": "MongoDB",
    "gcp": "GCP",
    "google cloud platform": "Google Cloud",
    "k8s": "Kubernetes",
    "ml": "Machine Learning",
    "dl": "Deep Learning",
    "sklearn": "Scikit-learn",
    "scikit learn": "Scikit-learn",
    "tf": "TensorFlow",
    "torch": "PyTorch",
    "excel": "Data Analysis",
    "dsa": "Data Structures",
    "data structures and algorithms": "Data Structures",
    "object oriented programming": "OOP",
    "restful": "REST",
    "ci cd": "CI/CD",
    "visual studio code": "VS Code",
}


def _norm(text: str) -> str:
    return " ".join(text.lower().split())


# ---------- File extraction helpers ----------

def _get_extension(filename: str) -> str:
    if not filename or "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[1].lower()


def _clean_text(text: str) -> str:
    cleaned_lines = []
    prev_blank = False
    for line in text.splitlines():
        stripped = line.strip()
        if stripped == "":
            if not prev_blank:
                cleaned_lines.append("")
            prev_blank = True
        else:
            cleaned_lines.append(" ".join(stripped.split()))
            prev_blank = False
    return "\n".join(cleaned_lines).strip()


def _extract_pdf(contents: bytes) -> str:
    pages_text = []
    with pdfplumber.open(io.BytesIO(contents)) as pdf:
        for page in pdf.pages:
            text = page.extract_text(layout=True) or page.extract_text() or ""
            pages_text.append(text)
    return "\n\n".join(pages_text)


def _extract_docx(contents: bytes) -> str:
    doc = Document(io.BytesIO(contents))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


async def extract_text_from_resume(file: UploadFile) -> str:
    if file is None or not file.filename:
        raise HTTPException(status_code=400, detail="No file provided.")

    ext = _get_extension(file.filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{ext or 'unknown'}'. Please upload a PDF or DOCX.",
        )

    contents = await file.read()

    MAX_SIZE = 10 * 1024 * 1024  # 10 MB
    if len(contents) > MAX_SIZE:
        raise HTTPException(
            status_code=400,
            detail=f"File is too large ({len(contents) / 1024 / 1024:.1f} MB). Maximum is 10 MB.",
        )

    if not contents:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    try:
        if ext == ".pdf":
            raw = _extract_pdf(contents)
        else:
            raw = _extract_docx(contents)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Failed to parse the file: {exc}")

    text = _clean_text(raw)
    if not text:
        raise HTTPException(
            status_code=400,
            detail="Could not extract any text. The file may be a scanned image PDF.",
        )
    return text


# ---------- Skill extraction ----------

def extract_skills(text: str, skill_list: list) -> list:
    if not text:
        return []

    doc = nlp(text)
    tokens = [_norm(tok.lemma_) for tok in doc if not tok.is_space]
    token_set = set(tokens)
    text_lower = _norm(text)

    found = set()
    for skill in skill_list:
        s_norm = _norm(skill)
        if " " in s_norm or "." in s_norm:
            if s_norm in text_lower:
                found.add(skill)
        else:
            if s_norm in token_set:
                found.add(skill)

    for alias, canonical in SKILL_ALIASES.items():
        a_norm = _norm(alias)
        if " " in a_norm:
            if a_norm in text_lower:
                found.add(canonical)
        else:
            if a_norm in token_set:
                found.add(canonical)

    return [s for s in skill_list if s in found]


# ---------- Semantic skill matching ----------

def _dedupe_by_similarity(skills: list, threshold: float = 0.85) -> list:
    if not skills:
        return []
    kept = []
    for skill in skills:
        vec = embed_skill(skill)
        duplicate = False
        for existing in kept:
            sim = util.cos_sim(vec, embed_skill(existing)).item()
            if sim >= threshold:
                duplicate = True
                break
        if not duplicate:
            kept.append(skill)
    return kept


def match_skills(resume_skills: list, job_skills: list, threshold: float = 0.6) -> dict:
    resume_skills = _dedupe_by_similarity(resume_skills)
    job_skills = _dedupe_by_similarity(job_skills)

    matched = []
    missing = []

    for job_skill in job_skills:
        job_vec = embed_skill(job_skill)
        best_sim = 0.0
        best_resume = None
        for r_skill in resume_skills:
            r_vec = embed_skill(r_skill)
            sim = util.cos_sim(job_vec, r_vec).item()
            if sim > best_sim:
                best_sim = sim
                best_resume = r_skill

        if best_sim >= threshold and best_resume is not None:
            matched.append({
                "job_skill": job_skill,
                "resume_skill": best_resume,
                "similarity": round(best_sim, 3),
            })
        else:
            missing.append(job_skill)

    matched_resume = {m["resume_skill"] for m in matched}
    extra = [s for s in resume_skills if s not in matched_resume]

    match_percentage = (
        round(len(matched) / len(job_skills) * 100, 1) if job_skills else 0.0
    )

    return {
        "matched_skills": matched,
        "missing_skills": missing,
        "extra_resume_skills": extra,
        "match_percentage": match_percentage,
    }


# ---------- Routes ----------

@app.get("/")
def root():
    return {"status": "ok", "service": "AI Skill Gap Analyzer API"}


# ============ AUTH ROUTES ============

@app.post("/signup")
def signup(req: SignupRequest, db: Session = Depends(get_db)):
    username = req.username.strip()
    email = req.email.strip().lower()

    if len(username) < 3:
        raise HTTPException(400, "Username must be at least 3 characters.")
    if "@" not in email:
        raise HTTPException(400, "Please enter a valid email address.")
    if len(req.password) < 6:
        raise HTTPException(400, "Password must be at least 6 characters.")

    existing = (
        db.query(User)
        .filter((User.username == username) | (User.email == email))
        .first()
    )
    if existing:
        raise HTTPException(400, "Username or email already registered.")

    user = User(
        username=username,
        email=email,
        password_hash=hash_password(req.password),
    )
    db.add(user)
    db.commit()
    db.refresh(user)

    return {"status": "ok", "user": user.to_dict()}


@app.post("/login")
def login(req: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.username == req.username.strip()).first()
    if not user or not verify_password(req.password, user.password_hash):
        raise HTTPException(401, "Invalid username or password.")

    return {"status": "ok", "user": user.to_dict()}


# ============ ANALYSIS ROUTE ============

@app.post("/analyze")
async def analyze(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
    db: Session = Depends(get_db),
):
    resume_text = await extract_text_from_resume(resume)

    resume_skills = extract_skills(resume_text, SKILL_LIST)
    job_skills = extract_skills(job_description, SKILL_LIST)

    matching = match_skills(resume_skills, job_skills)

    # --- Fuzzy / string-match features for the ML model ---
    resume_skill_str = " ".join(resume_skills)
    job_skill_str = " ".join(job_skills)

    skill_string_match_score = (
        fuzz.token_set_ratio(resume_skill_str, job_skill_str)
        if resume_skills and job_skills else 0.0
    )

    if job_skills and resume_skills:
        best_scores = []
        for j_skill in job_skills:
            best = max(
                fuzz.token_set_ratio(j_skill, r_skill) for r_skill in resume_skills
            )
            best_scores.append(best)
        fuzzy_match_score = sum(best_scores) / len(best_scores)
    else:
        fuzzy_match_score = 0.0

    scoring = analyze_with_rules(
        match_percentage=matching["match_percentage"],
        matched_count=len(matching["matched_skills"]),
        resume_skill_count=len(resume_skills),
        resume_text=resume_text,
        resume_skills=resume_skills,
        job_skills=job_skills,
        job_text=job_description,
        skill_string_match_score=skill_string_match_score,
        fuzzy_match_score=fuzzy_match_score,
    )

    roadmap = generate_roadmap(matching["missing_skills"])

    # --- Persist to database ---
    try:
        record = Analysis(
            resume_filename=resume.filename,
            job_description=job_description,
            match_percentage=matching["match_percentage"],
            resume_score=scoring["resume_score"],
            interview_probability=scoring["interview_probability"],
            project_count=scoring["features"].get("project_count", 0),
            matched_skills=matching["matched_skills"],
            missing_skills=matching["missing_skills"],
            extra_resume_skills=matching["extra_resume_skills"],
            roadmap=roadmap,
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        analysis_id = record.id
    except Exception as exc:
        db.rollback()
        print(f"[db] failed to save analysis: {exc}")
        analysis_id = None

    return {
        "status": "matched",
        "analysis_id": analysis_id,
        "resume_filename": resume.filename,
        "resume_text_length": len(resume_text),
        "job_description_chars": len(job_description),
        "resume_skills": resume_skills,
        "job_skills": job_skills,
        "matched_skills": matching["matched_skills"],
        "missing_skills": matching["missing_skills"],
        "extra_resume_skills": matching["extra_resume_skills"],
        "match_percentage": matching["match_percentage"],
        "resume_score": scoring["resume_score"],
        "interview_probability": scoring["interview_probability"],
        "features": scoring["features"],
        "model_used": scoring["model_used"],
        "roadmap": roadmap,
    }


# ============ HISTORY ROUTES ============

@app.get("/history")
def history(limit: int = 20, db: Session = Depends(get_db)):
    rows = (
        db.query(Analysis)
        .order_by(Analysis.timestamp.desc())
        .limit(limit)
        .all()
    )
    return {"count": len(rows), "items": [r.to_dict() for r in rows]}


@app.get("/history/{analysis_id}")
def history_detail(analysis_id: int, db: Session = Depends(get_db)):
    row = db.query(Analysis).filter(Analysis.id == analysis_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Analysis not found")
    return row.to_dict()


@app.delete("/history/{analysis_id}")
def delete_analysis(analysis_id: int, db: Session = Depends(get_db)):
    row = db.query(Analysis).filter(Analysis.id == analysis_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Analysis not found")
    db.delete(row)
    db.commit()
    return {"status": "deleted", "id": analysis_id}